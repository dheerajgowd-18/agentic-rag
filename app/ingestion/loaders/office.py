import os
import logfire
import docx as python_docx
from pptx import Presentation


def parse_office(file_path: str) -> str:
    """
    Parses Office documents (.docx, .pptx) extracting both paragraphs and structured tables.
    Direct parsing avoids unstructured layout model hangs on Windows.
    """
    with logfire.span("📄 Office Document Parsing", filename=file_path):
        try:
            ext = os.path.splitext(file_path)[1].lower()
            texts = []

            if ext == ".docx":
                doc = python_docx.Document(file_path)
                # 1. Paragraphs
                for p in doc.paragraphs:
                    if p.text.strip():
                        texts.append(p.text.strip())

                # 2. Tables
                for table in doc.tables:
                    for row in table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            # Deduplicate repeated merged cell text in the same row
                            unique_cells = []
                            for c in row_cells:
                                if not unique_cells or c != unique_cells[-1]:
                                    unique_cells.append(c)
                            texts.append(" | ".join(unique_cells))

                full_text = "\n\n".join(texts)

            elif ext == ".pptx":
                prs = Presentation(file_path)
                for slide in prs.slides:
                    for shape in slide.shapes:
                        # 1. Text shapes
                        if hasattr(shape, "text") and shape.text.strip():
                            texts.append(shape.text.strip())

                        # 2. Table shapes
                        if getattr(shape, "has_table", False):
                            for row in shape.table.rows:
                                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                                if row_cells:
                                    texts.append(" | ".join(row_cells))

                full_text = "\n\n".join(texts)
            else:
                full_text = ""

            if not full_text.strip():
                logfire.warning(f"⚠️ Empty text extracted for {file_path}")
            else:
                logfire.info(f"✅ Successfully parsed {len(full_text)} characters from {file_path}")

            return full_text
        except Exception as e:
            logfire.error(f"❌ Office Parse Failed for {file_path}: {e}")
            raise e
